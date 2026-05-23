import os
import json
from dataclasses import dataclass
from datetime import datetime
from xml.sax.saxutils import escape
from typing import Any, Dict, List, Optional
from sqlalchemy.orm import Session

from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
import matplotlib.patches as mpatches

from backend.app.config import get_settings
from backend.app import models

settings = get_settings()


@dataclass
class ReportTheme:
    id: str
    name: str
    definition: str = ""
    theme_type: str = "overarching"
    parent_id: Optional[str] = None
    parent_name: Optional[str] = None


class ReportService:
    def __init__(self, db: Session):
        self.db = db
        self.template_dir = os.path.join(os.path.dirname(__file__), "..", "templates")

    def _phase_structured_data(self, project_id: str, phase_number: int) -> Dict[str, Any]:
        phase_state = (
            self.db.query(models.PhaseState)
            .filter(models.PhaseState.project_id == project_id)
            .filter(models.PhaseState.phase_number == phase_number)
            .first()
        )
        return phase_state.structured_data if phase_state and phase_state.structured_data else {}

    def _themes_from_structured_data(self, data: Dict[str, Any]) -> List[ReportTheme]:
        raw_themes = (
            data.get("final_themes")
            or data.get("refined_themes")
            or data.get("candidate_themes")
            or []
        )

        themes: List[ReportTheme] = []

        for index, raw in enumerate(raw_themes):
            if isinstance(raw, str):
                raw = {"name": raw}
            if not isinstance(raw, dict) or not raw.get("name"):
                continue

            theme_name = str(raw["name"])
            parent_name = raw.get("parent_name") or raw.get("parent")
            theme_type = raw.get("theme_type") or raw.get("type") or ("sub" if parent_name else "overarching")
            theme = ReportTheme(
                id=f"theme-{index}",
                name=theme_name,
                definition=str(raw.get("definition") or ""),
                theme_type=theme_type,
                parent_name=parent_name,
            )
            themes.append(theme)

            for sub_index, sub_theme in enumerate(raw.get("sub_themes") or []):
                if isinstance(sub_theme, str):
                    sub_theme = {"name": sub_theme}
                if not isinstance(sub_theme, dict) or not sub_theme.get("name"):
                    continue
                themes.append(
                    ReportTheme(
                        id=f"theme-{index}-sub-{sub_index}",
                        name=str(sub_theme["name"]),
                        definition=str(sub_theme.get("definition") or ""),
                        theme_type="sub",
                        parent_name=theme_name,
                    )
                )

        name_to_id = {theme.name: theme.id for theme in themes}
        for theme in themes:
            if theme.parent_name and not theme.parent_id:
                theme.parent_id = name_to_id.get(theme.parent_name)

        return themes

    def _get_themes(self, project: models.Project) -> List[Any]:
        db_themes = (
            self.db.query(models.Theme)
            .filter(models.Theme.project_id == project.id)
            .all()
        )
        if db_themes:
            return db_themes

        for phase_number in (6, 5, 4):
            themes = self._themes_from_structured_data(
                self._phase_structured_data(project.id, phase_number)
            )
            if themes:
                return themes

        return []

    def _get_project_data(self, project: models.Project) -> Dict[str, Any]:
        """Gather all project data for report generation."""
        themes = self._get_themes(project)

        # Get final map
        final_map = (
            self.db.query(models.ThematicMap)
            .filter(models.ThematicMap.project_id == project.id)
            .filter(models.ThematicMap.phase_label == "final")
            .first()
        )

        # Get phase 7 state for report text
        phase7_state = (
            self.db.query(models.PhaseState)
            .filter(models.PhaseState.project_id == project.id)
            .filter(models.PhaseState.phase_number == 7)
            .first()
        )

        report_text = ""
        extracts_for_report = []
        if phase7_state and phase7_state.structured_data:
            report_text = phase7_state.structured_data.get("report_text", "")
            extracts_for_report = phase7_state.structured_data.get("extracts_for_report", [])

        return {
            "project": project,
            "themes": themes,
            "final_map_path": final_map.image_path if final_map else None,
            "report_text": report_text,
            "extracts_for_report": extracts_for_report,
            "analytic_decisions": project.analytic_decisions or {},
            "research_question": project.research_question or "",
            "source_count": len(project.sources),
        }

    def generate_thematic_map_image(
        self,
        project_id: str,
        themes: List[Any],
        phase_label: str = "final",
    ) -> str:
        """Generate a thematic map PNG using matplotlib, adapted from references/thematic-map.md."""
        # Organize themes
        overarching = [t for t in themes if t.theme_type == "overarching"]
        subs = [t for t in themes if t.theme_type == "sub"]

        fig, ax = plt.subplots(figsize=(14, 10), dpi=300)
        ax.set_xlim(0, 14)
        ax.set_ylim(0, 10)
        ax.axis("off")

        THEME_FACE = "#E8E8E8"
        SUBTHEME_FACE = "#FFFFFF"
        EDGE = "#333333"
        FONT = {"family": "serif", "fontsize": 11}

        def draw_theme(x, y, label, width=3.0, height=1.2):
            ellipse = mpatches.Ellipse(
                (x, y), width, height,
                facecolor=THEME_FACE, edgecolor=EDGE, linewidth=1.5
            )
            ax.add_patch(ellipse)
            ax.text(x, y, label, ha="center", va="center", fontweight="bold", wrap=True, **FONT)
            return (x, y)

        def draw_subtheme(x, y, label, width=2.4, height=0.7):
            rect = mpatches.FancyBboxPatch(
                (x - width / 2, y - height / 2),
                width, height,
                boxstyle="round,pad=0.05",
                facecolor=SUBTHEME_FACE,
                edgecolor=EDGE, linewidth=1.0
            )
            ax.add_patch(rect)
            ax.text(x, y, label, ha="center", va="center", wrap=True, fontsize=9, family="serif")
            return (x, y)

        def connect(x1, y1, x2, y2):
            ax.plot([x1, x2], [y1, y2], color=EDGE, linewidth=0.8, zorder=0)

        # Simple layout: distribute overarching themes across top
        n_over = len(overarching)
        if n_over == 0:
            n_over = 1
            overarching = [None]

        positions = []
        for i, theme in enumerate(overarching):
            x = 2 + (i * (10 / max(n_over - 1, 1))) if n_over > 1 else 7
            y = 8.5
            if theme:
                pos = draw_theme(x, y, theme.name)
                positions.append((theme, pos))
            else:
                positions.append((None, (x, y)))

        # Place sub-themes under their parents
        for theme, (tx, ty) in positions:
            if not theme:
                continue
            theme_subs = [s for s in subs if s.parent_id == theme.id]
            n_subs = len(theme_subs)
            for j, sub in enumerate(theme_subs):
                sx = tx - 1.5 + (j * (3 / max(n_subs - 1, 1))) if n_subs > 1 else tx
                sy = 6.0 - (j % 2) * 0.3  # slight stagger
                draw_subtheme(sx, sy, sub.name)
                connect(tx, ty - 0.6, sx, sy + 0.35)

        plt.tight_layout()

        out_path = os.path.join(settings.maps_dir, f"{project_id}_{phase_label}_map.png")
        plt.savefig(out_path, dpi=300, bbox_inches="tight", facecolor="white")
        plt.close()

        # Save to DB
        existing = (
            self.db.query(models.ThematicMap)
            .filter(models.ThematicMap.project_id == project_id)
            .filter(models.ThematicMap.phase_label == phase_label)
            .first()
        )
        if existing:
            existing.image_path = out_path
        else:
            db_map = models.ThematicMap(
                project_id=project_id,
                phase_label=phase_label,
                image_path=out_path,
                caption=f"{phase_label.capitalize()} thematic map",
            )
            self.db.add(db_map)
        self.db.commit()

        return out_path

    def generate_report(self, project: models.Project, format: str = "pdf") -> str:
        """Generate a downloadable report (PDF or DOCX)."""
        data = self._get_project_data(project)

        # Ensure thematic map exists
        existing_map = (
            self.db.query(models.ThematicMap)
            .filter(models.ThematicMap.project_id == project.id)
            .filter(models.ThematicMap.phase_label == "final")
            .first()
        )
        if data["themes"] and (not existing_map or not os.path.exists(existing_map.image_path)):
            self.generate_thematic_map_image(project.id, data["themes"], phase_label="final")
            data = self._get_project_data(project)

        if format == "pdf":
            return self._generate_pdf(data)
        elif format == "docx":
            return self._generate_docx(data)
        else:
            raise ValueError(f"Unsupported format: {format}")

    def _generate_pdf(self, data: Dict[str, Any]) -> str:
        """Generate a PDF report with ReportLab."""
        from reportlab.lib import colors
        from reportlab.lib.enums import TA_CENTER
        from reportlab.lib.pagesizes import A4
        from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
        from reportlab.lib.units import inch
        from reportlab.lib.utils import ImageReader
        from reportlab.platypus import Image, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

        project = data["project"]
        out_path = os.path.join(settings.reports_dir, f"{project.id}_report.pdf")

        styles = getSampleStyleSheet()
        title_style = ParagraphStyle(
            "ReportTitle",
            parent=styles["Title"],
            alignment=TA_CENTER,
            textColor=colors.HexColor("#4f46e5"),
            spaceAfter=12,
        )
        heading_style = ParagraphStyle(
            "ReportHeading",
            parent=styles["Heading2"],
            textColor=colors.HexColor("#4338ca"),
            spaceBefore=14,
            spaceAfter=8,
        )
        body_style = ParagraphStyle(
            "ReportBody",
            parent=styles["BodyText"],
            leading=14,
            spaceAfter=8,
        )
        small_style = ParagraphStyle(
            "ReportSmall",
            parent=styles["BodyText"],
            fontSize=9,
            leading=12,
            textColor=colors.HexColor("#64748b"),
            alignment=TA_CENTER,
            spaceAfter=14,
        )

        def text(value: Any) -> str:
            return escape(str(value or "")).replace("\n", "<br/>")

        story = [
            Paragraph(text(project.title), title_style),
            Paragraph(
                f"Thematic Analysis Report | Braun &amp; Clarke (2006) Six-Phase Framework | Generated {datetime.utcnow().strftime('%Y-%m-%d')}",
                small_style,
            ),
            Paragraph("Research Question", heading_style),
            Paragraph(text(data.get("research_question") or "Not specified"), body_style),
            Paragraph("Analytic Decisions", heading_style),
            Paragraph(text(json.dumps(data.get("analytic_decisions", {}), indent=2)), body_style),
            Paragraph("Themes", heading_style),
        ]

        table_data = [[
            Paragraph("<b>Name</b>", body_style),
            Paragraph("<b>Definition</b>", body_style),
            Paragraph("<b>Type</b>", body_style),
        ]]
        for theme in data.get("themes", []):
            table_data.append([
                Paragraph(text(theme.name), body_style),
                Paragraph(text(theme.definition), body_style),
                Paragraph(text(theme.theme_type), body_style),
            ])

        if len(table_data) == 1:
            story.append(Paragraph("No themes generated yet.", body_style))
        else:
            table = Table(table_data, colWidths=[1.6 * inch, 3.6 * inch, 1.0 * inch], repeatRows=1)
            table.setStyle(TableStyle([
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#eef2ff")),
                ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#cbd5e1")),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("LEFTPADDING", (0, 0), (-1, -1), 6),
                ("RIGHTPADDING", (0, 0), (-1, -1), 6),
                ("TOPPADDING", (0, 0), (-1, -1), 6),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
            ]))
            story.extend([table, Spacer(1, 8)])

        map_path = data.get("final_map_path")
        if map_path and os.path.exists(map_path):
            story.append(Paragraph("Thematic Map", heading_style))
            image_reader = ImageReader(map_path)
            image_width, image_height = image_reader.getSize()
            display_width = 6.2 * inch
            display_height = display_width * (image_height / image_width)
            story.append(Image(map_path, width=display_width, height=display_height))
            story.append(Paragraph("Figure 1. Final thematic map showing overarching themes and sub-themes.", small_style))

        story.append(Paragraph("Analysis Report", heading_style))
        report_text = data.get("report_text") or "No report text generated yet."
        for paragraph in str(report_text).split("\n\n"):
            if paragraph.strip():
                story.append(Paragraph(text(paragraph.strip()), body_style))

        doc = SimpleDocTemplate(
            out_path,
            pagesize=A4,
            rightMargin=40,
            leftMargin=40,
            topMargin=40,
            bottomMargin=40,
        )
        doc.build(story)
        return out_path

    def _generate_docx(self, data: Dict[str, Any]) -> str:
        """Generate DOCX manuscript."""
        project = data["project"]
        out_path = os.path.join(settings.reports_dir, f"{project.id}_report.docx")

        doc = Document()

        # Title
        title = doc.add_heading(project.title, level=0)
        title.alignment = WD_ALIGN_PARAGRAPH.CENTER

        # Meta
        meta = doc.add_paragraph()
        meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
        meta_run = meta.add_run(f"Thematic Analysis Report | Generated {datetime.utcnow().strftime('%Y-%m-%d')}")
        meta_run.font.size = Pt(10)
        meta_run.font.color.rgb = RGBColor(0x64, 0x70, 0x8B)

        doc.add_paragraph()

        # Research question
        doc.add_heading("Research Question", level=1)
        doc.add_paragraph(data.get("research_question", "Not specified"))

        # Analytic decisions
        doc.add_heading("Analytic Decisions", level=1)
        decisions = data.get("analytic_decisions", {})
        for key, value in decisions.items():
            p = doc.add_paragraph(style="List Bullet")
            p.add_run(f"{key}: ").bold = True
            p.add_run(str(value))

        # Themes
        doc.add_heading("Themes", level=1)
        themes = data.get("themes", [])
        for t in themes:
            doc.add_heading(t.name, level=2)
            doc.add_paragraph(t.definition or "")

        # Thematic map
        map_path = data.get("final_map_path")
        if map_path and os.path.exists(map_path):
            doc.add_heading("Thematic Map", level=1)
            doc.add_picture(map_path, width=Inches(6))
            caption = doc.add_paragraph("Figure 1. Final thematic map.")
            caption.alignment = WD_ALIGN_PARAGRAPH.CENTER

        # Report text
        report_text = data.get("report_text", "")
        if report_text:
            doc.add_heading("Findings", level=1)
            doc.add_paragraph(report_text)

        doc.save(out_path)
        return out_path
